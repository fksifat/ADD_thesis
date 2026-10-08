import os
import sys
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import latex2mathml.converter
import mathml2omml

def latex_to_omml_element(latex_str):
    """Converts a LaTeX math string into a native Word OMML XML element."""
    try:
        # Pre-clean known edge cases for latex2mathml
        cleaned = latex_str.strip()
        cleaned = re.sub(r'\\bar\{([^}]+)\}', r'\\overline{\1}', cleaned)
        cleaned = re.sub(r'\\hat\{([^}]+)\}', r'\\widehat{\1}', cleaned)
        cleaned = cleaned.replace(r'\parallel', '||')
        cleaned = cleaned.replace(r'\text{ \AA}', r'\text{ Å}')
        cleaned = cleaned.replace(r'\AA', r'\text{Å}')
        
        mathml = latex2mathml.converter.convert(cleaned)
        omml = mathml2omml.convert(mathml)
        
        # Ensure proper math namespace declarations
        if 'xmlns:m=' not in omml:
            omml = omml.replace('<m:oMath>', f'<m:oMath {nsdecls("m")}>')
            omml = omml.replace('<m:oMathPara>', f'<m:oMathPara {nsdecls("m")}>')
        
        return parse_xml(omml)
    except Exception as e:
        print(f"OMML conversion fallback for '{latex_str[:30]}...': {e}")
        return None

def format_math_unicode(latex_str):
    """Fallback: converts LaTeX math symbols to clean Unicode mathematical typography."""
    text = latex_str.strip()
    replacements = [
        (r'\mathbf{H}', '𝐇'), (r'\mathbf{h}', '𝐡'), (r'\mathbf{e}', '𝐞'),
        (r'\mathbf{m}', '𝐦'), (r'\mathbf{r}', '𝐫'), (r'\mathbf{s}', '𝐬'),
        (r'\mathbf{u}', '𝐮'), (r'\mathbf{U}', '𝐔'), (r'\mathbf{W}', '𝐖'),
        (r'\mathbf{b}', '𝐛'), (r'\mathbf{A}', '𝐀'), (r'\mathbf{B}', '𝐁'),
        (r'\mathbf{g}', '𝐠'), (r'\mathbf{T}', '𝐓'), (r'\mathbf{Z}', '𝐙'),
        (r'\mathbf{x}', '𝐱'), (r'\mathbf{y}', '𝐲'), (r'\mathbf{z}', '𝐳'),
        (r'\mathbf{R}', '𝐑'), (r'\boldsymbol{\alpha}', '𝜶'),
        (r'\mathbb{R}', 'ℝ'), (r'\Delta', 'Δ'), (r'\mathcal{L}', 'ℒ'),
        (r'\mathcal{N}', '𝒩'), (r'\mathcal{M}', 'ℳ'), (r'\in', '∈'),
        (r'\sum', '∑'), (r'\dots', '…'), (r'\times', '×'),
        (r'\tau', 'τ'), (r'\gamma', 'γ'), (r'\mu', 'μ'), (r'\sigma', 'σ'),
        (r'\phi', 'ϕ'), (r'\alpha', 'α'), (r'\hat', '^'), (r'\bar', '¯'),
        (r'\parallel', '||'), (r'\le', '≤'), (r'\ge', '≥'), (r'\neq', '≠'),
        (r'\to', '→'), (r'\quad', '  '), (r'\,', ' '),
        (r'\left(', '('), (r'\right)', ')'), (r'\left[', '['), (r'\right]', ']'),
        (r'\left\{', '{'), (r'\right\}', '}'), (r'\text{', ''),
        (r'\mathbf{', ''), (r'\mathbb{', ''), (r'\mathcal{', ''),
        (r'}', '')
    ]
    for old, new in replacements:
        text = text.replace(old, new)
    return text

def create_docx(md_path, docx_path):
    doc = docx.Document()
    
    # 1-inch margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header / Footer
        footer_p = section.footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = footer_p.add_run("TriBioNode v2 Research Manuscript | Page ")
        f_run.font.name = "Calibri"
        f_run.font.size = Pt(9)
        f_run.font.color.rgb = RGBColor(128, 128, 128)

    # Base style
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = RGBColor(30, 41, 59) # Slate 800
    style_normal.paragraph_format.line_spacing = 1.18
    style_normal.paragraph_format.space_after = Pt(6)

    def set_cell_shading(cell, color_hex):
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shd)

    def set_cell_margins(cell, top=100, bottom=100, left=130, right=130):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def set_table_borders(table):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="12" w:space="0" w:color="1E3A8A"/>'
            f'<w:bottom w:val="single" w:sz="12" w:space="0" w:color="1E3A8A"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
            f'<w:insideV w:val="none"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    def add_formatted_text(p, text, is_bold_default=False, is_italic_default=False, font_size=11, font_color=None):
        parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\$.*?\$)', text)
        for part in parts:
            if not part:
                continue
            
            # Inline LaTeX math ($...$)
            if part.startswith('$') and part.endswith('$') and len(part) >= 2:
                math_code = part[1:-1].strip()
                omml_elem = latex_to_omml_element(math_code)
                if omml_elem is not None:
                    p._p.append(omml_elem)
                else:
                    # Fallback to Cambria Math unicode formatted text
                    run = p.add_run(format_math_unicode(math_code))
                    run.font.name = 'Cambria Math'
                    run.font.size = Pt(font_size)
                    run.italic = True
                continue

            run = p.add_run()
            run.font.name = 'Calibri'
            run.font.size = Pt(font_size)
            if font_color:
                run.font.color.rgb = font_color
            
            if part.startswith('**') and part.endswith('**') and len(part) >= 4:
                run.text = part[2:-2]
                run.bold = True
                run.italic = is_italic_default
            elif part.startswith('*') and part.endswith('*') and len(part) >= 2:
                run.text = part[1:-1]
                run.bold = is_bold_default
                run.italic = True
            elif part.startswith('`') and part.endswith('`') and len(part) >= 2:
                run.text = part[1:-1]
                run.font.name = 'Consolas'
                run.font.size = Pt(font_size - 1)
                run.font.color.rgb = RGBColor(180, 50, 10)
            else:
                run.text = part
                run.bold = is_bold_default
                run.italic = is_italic_default

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    i = 0
    table_lines = []

    while i < len(lines):
        line = lines[i].rstrip('\n')
        stripped = line.strip()

        # Handle Tables
        if stripped.startswith('|') and stripped.endswith('|'):
            table_lines.append(stripped)
            i += 1
            while i < len(lines) and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            
            rows_data = []
            for tl in table_lines:
                if re.match(r'^\|[\s\-:\+\|]+$', tl):
                    continue
                cols = [c.strip() for c in tl.split('|')[1:-1]]
                rows_data.append(cols)
            
            if rows_data:
                num_rows = len(rows_data)
                num_cols = len(rows_data[0])
                table = doc.add_table(rows=num_rows, cols=num_cols)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                set_table_borders(table)

                for r_idx, row in enumerate(rows_data):
                    table_row = table.rows[r_idx]
                    is_header = (r_idx == 0)
                    for c_idx in range(min(num_cols, len(row))):
                        cell = table_row.cells[c_idx]
                        set_cell_margins(cell, top=100, bottom=100, left=130, right=130)
                        if is_header:
                            set_cell_shading(cell, 'F1F5F9')
                        elif r_idx % 2 == 1:
                            set_cell_shading(cell, 'F8FAFC')
                        
                        cell_p = cell.paragraphs[0]
                        cell_p.paragraph_format.space_after = Pt(2)
                        cell_p.paragraph_format.space_before = Pt(2)
                        cell_p.paragraph_format.line_spacing = 1.05
                        
                        cell_text = row[c_idx]
                        if c_idx > 0 and (re.search(r'\d', cell_text) or is_header):
                            cell_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        else:
                            cell_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                            
                        add_formatted_text(cell_p, cell_text, is_bold_default=is_header, font_size=9.5)
                
                post_p = doc.add_paragraph()
                post_p.paragraph_format.space_after = Pt(6)
                post_p.paragraph_format.space_before = Pt(0)
            
            table_lines = []
            continue

        # Document Title
        if stripped.startswith('# TriBioNode') or (stripped.startswith('# ') and i < 5):
            title_text = stripped.lstrip('#').strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(10)
            run = p.add_run(title_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(18)
            run.bold = True
            run.font.color.rgb = RGBColor(30, 58, 138)
            i += 1
            continue

        # Author / Affiliation
        if (stripped.startswith('**Farhan') or stripped.startswith('*Department') or stripped.startswith('**Authors**:')) and i < 15:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            add_formatted_text(p, stripped, font_size=11, font_color=RGBColor(71, 85, 105))
            i += 1
            continue

        # Divider
        if stripped == '---':
            i += 1
            continue

        # Abstract
        if stripped.startswith('## Abstract'):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run('ABSTRACT')
            run.font.name = 'Calibri'
            run.font.size = Pt(12)
            run.bold = True
            run.font.color.rgb = RGBColor(30, 58, 138)
            
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.cell(0, 0)
            set_cell_shading(cell, 'F8FAFC')
            set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
            
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="1E3A8A"/><w:top w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>')
            tcPr.append(tcBorders)
            
            i += 1
            abs_lines = []
            while i < len(lines) and not lines[i].strip().startswith('#') and not lines[i].strip() == '---':
                if lines[i].strip():
                    abs_lines.append(lines[i].strip())
                i += 1
            
            abs_text = ' '.join(abs_lines)
            abs_p = cell.paragraphs[0]
            abs_p.paragraph_format.line_spacing = 1.15
            abs_p.paragraph_format.space_after = Pt(2)
            add_formatted_text(abs_p, abs_text, font_size=10.5, font_color=RGBColor(51, 65, 85))
            
            doc.add_paragraph().paragraph_format.space_after = Pt(8)
            continue

        # Heading 1
        if stripped.startswith('# '):
            h_text = stripped.lstrip('#').strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(13.5)
            run.bold = True
            run.font.color.rgb = RGBColor(30, 58, 138)
            i += 1
            continue

        # Heading 2
        if stripped.startswith('## '):
            h_text = stripped.lstrip('#').strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(13)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(12)
            run.bold = True
            run.font.color.rgb = RGBColor(37, 99, 235)
            i += 1
            continue

        # Heading 3
        if stripped.startswith('### '):
            h_text = stripped.lstrip('#').strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(11)
            run.bold = True
            run.font.color.rgb = RGBColor(30, 41, 59)
            i += 1
            continue

        # Images
        img_match = re.match(r'^!\[(.*?)\]\((.*?)\)$', stripped)
        if img_match:
            raw_path = img_match.group(2)
            if raw_path.startswith('../'):
                resolved_path = os.path.normpath(os.path.join(os.path.dirname(md_path), raw_path))
            elif raw_path.startswith('file://'):
                resolved_path = raw_path.replace('file://', '')
            else:
                resolved_path = os.path.normpath(os.path.join(os.path.dirname(md_path), raw_path))
            
            if os.path.exists(resolved_path):
                img_p = doc.add_paragraph()
                img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                img_p.paragraph_format.space_before = Pt(10)
                img_p.paragraph_format.space_after = Pt(2)
                img_p.paragraph_format.keep_with_next = True
                
                try:
                    img_p.add_run().add_picture(resolved_path, width=Inches(5.8))
                except Exception as e:
                    print(f"Image warning: {e}")
            
            i += 1
            if i < len(lines) and lines[i].strip().startswith('*Figure'):
                cap_text = lines[i].strip().strip('*')
                cap_p = doc.add_paragraph()
                cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap_p.paragraph_format.space_before = Pt(2)
                cap_p.paragraph_format.space_after = Pt(10)
                run = cap_p.add_run(cap_text)
                run.font.name = 'Calibri'
                run.font.size = Pt(9.5)
                run.italic = True
                run.font.color.rgb = RGBColor(100, 116, 139)
                i += 1
            continue

        # Standalone Block Equations ($$ ... $$)
        if stripped.startswith('$$') and stripped.endswith('$$') and len(stripped) > 4:
            eq_code = stripped[2:-2].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            
            omml_elem = latex_to_omml_element(eq_code)
            if omml_elem is not None:
                p._p.append(omml_elem)
            else:
                run = p.add_run(format_math_unicode(eq_code))
                run.font.name = 'Cambria Math'
                run.font.size = Pt(11)
                run.italic = True
                run.font.color.rgb = RGBColor(15, 23, 42)
            i += 1
            continue

        # Bullet lists
        if stripped.startswith('- ') or stripped.startswith('* '):
            item_text = stripped[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_text(p, item_text)
            i += 1
            continue

        # Numbered lists
        num_match = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if num_match and not stripped.startswith('#'):
            item_text = num_match.group(2)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_text(p, item_text)
            i += 1
            continue

        # Regular Paragraph
        if stripped:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = 1.18
            add_formatted_text(p, stripped)

        i += 1

    doc.save(docx_path)
    print(f"Docx successfully built with OMML math equations at {docx_path} ({os.path.getsize(docx_path)} bytes)")

if __name__ == '__main__':
    create_docx('reports/TRIBIONODE_V2_PAPER_DRAFT.md', 'reports/TriBioNode_v2_Paper_Draft.docx')
